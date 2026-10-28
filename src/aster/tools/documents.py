from __future__ import annotations

import json
import sys
from pathlib import Path

from jsonschema import validate

from aster.domain import ToolContext, ToolResult, ToolSpec
from aster.platform.processes import ProcessSupervisor
from aster.tools.files import TEXT, artifact, bounded_result, failure, read_path, schema


def extract_document(path: Path, *, max_units=1000, max_chars=2_000_000) -> dict:
    """Runs in a disposable managed process; values retain source coordinates."""
    units, warnings, used = [], [], 0

    def add(location, text):
        nonlocal used
        if len(units) >= max_units or used >= max_chars:
            if "Extraction budget reached" not in warnings:
                warnings.append("Extraction budget reached")
            return False
        text = str(text)
        available = max_chars - used
        units.append({"location": location, "text": text[:available]})
        used += min(len(text), available)
        if len(text) > available:
            warnings.append("Extraction budget reached")
        return True

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF requires a decrypted input copy")
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                warnings.append(f"Page {number} has no extractable text; OCR may be required")
            if not add({"page": number}, text):
                break
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(path)
        paragraph, table = 0, 0
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                paragraph += 1
                if not add({"paragraph": paragraph, "style": block.style.name if block.style else None}, block.text):
                    break
            elif isinstance(block, Table):
                table += 1
                for row_number, row in enumerate(block.rows, 1):
                    if not add({"table": table, "row": row_number}, "\t".join(cell.text for cell in row.cells)):
                        break
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            warnings.append("Formula expressions are read as stored; no formula recalculation performed")
            for worksheet in workbook:
                for row in worksheet.iter_rows():
                    if len(units) >= max_units or used >= max_chars:
                        add({}, "")
                        break
                    for cell in row:
                        if cell.value is not None and not add({"sheet": worksheet.title, "cell": cell.coordinate}, cell.value):
                            break
        finally:
            workbook.close()
    elif suffix == ".pptx":
        from pptx import Presentation
        presentation = Presentation(path)
        warnings.append("SmartArt, images, charts and embedded objects may not have extractable text")
        for number, slide in enumerate(presentation.slides, 1):
            for shape_number, shape in enumerate(slide.shapes, 1):
                if shape.has_text_frame and not add({"slide": number, "shape": shape_number}, shape.text_frame.text):
                    break
                if shape.has_table:
                    for row_number, row in enumerate(shape.table.rows, 1):
                        if not add({"slide": number, "shape": shape_number, "row": row_number},
                                   "\t".join(cell.text for cell in row.cells)):
                            break
            if slide.has_notes_slide:
                text_frame = slide.notes_slide.notes_text_frame
                if text_frame:
                    add({"slide": number, "notes": True}, text_frame.text)
    else:
        raise ValueError("Supported document formats: PDF, DOCX, XLSX, PPTX")
    return {"source": str(path.resolve()), "format": suffix[1:], "units": units,
            "warnings": warnings, "method": "native-python-text-extraction"}


class DocumentTool:
    def __init__(self, supervisor: ProcessSupervisor, extension: str | None = None):
        self.supervisor, self.extension = supervisor, extension
        self.spec = ToolSpec(f"{extension}_to_text" if extension else "document_extract",
                             "Extract document content with source locations; no OCR or layout editing",
                             schema({"path": TEXT, "max_units": {"type": "integer", "minimum": 1, "maximum": 10000},
                                     "max_chars": {"type": "integer", "minimum": 1, "maximum": 5_000_000},
                                     "timeout": {"type": "number", "minimum": 1, "maximum": 300}}, ["path"]),
                             concurrency_safe=True, replay_safe=True, timeout_seconds=315)

    async def execute(self, arguments: dict, context: ToolContext) -> ToolResult:
        try:
            validate(arguments, self.spec.parameters)
            path = read_path(arguments["path"], context)
            if self.extension and path.suffix.lower() != "." + self.extension:
                raise ValueError(f"Expected a .{self.extension} file")
            if path.stat().st_size > 100_000_000:
                raise ValueError("Document exceeds 100 MB input budget")
            result = await self.supervisor.run(
                [sys.executable, "-m", "aster.tools.documents", str(path), str(arguments.get("max_units", 1000)),
                 str(arguments.get("max_chars", 2_000_000))], cwd=context.cwd,
                artifact_dir=context.artifact_dir, cancel=context.cancel, timeout=arguments.get("timeout", 120),
                env={"PYTHONIOENCODING": "utf-8"})
            if result.cancelled or result.timed_out or result.exit_code != 0 or result.error:
                category = "cancelled" if result.cancelled else "timeout" if result.timed_out else "document"
                return ToolResult(f"Document extraction failed: {result.error or result.stderr or category}",
                                  True, category, artifacts=result.artifacts)
            output = next(item for item in result.artifacts if item["stream"] == "stdout")
            payload = json.loads(Path(output["path"]).read_text(encoding="utf-8"))
            text = "\n\n".join(f"{json.dumps(unit['location'], ensure_ascii=False)}\n{unit['text']}"
                               for unit in payload["units"])
            if payload["warnings"]:
                text += "\n\nWarnings: " + "; ".join(payload["warnings"])
            answer = bounded_result(text, context, data={key: value for key, value in payload.items() if key != "units"})
            answer.artifacts.append(artifact(Path(output["path"]), context, "application/json", kind="document-extraction"))
            return answer
        except Exception as exc:  # noqa: BLE001 - tool boundary serializes extraction failures
            return failure(exc)


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        print(json.dumps(extract_document(Path(sys.argv[1]), max_units=int(sys.argv[2]), max_chars=int(sys.argv[3])),
                         ensure_ascii=False))
    except Exception as error:  # noqa: BLE001 - isolated worker's terminal error report
        print(f"{type(error).__name__}: {error}", file=sys.stderr)
        raise SystemExit(1)
