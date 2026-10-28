"""Explicitly budgeted manual acceptance driver. Never collected by pytest.

Run with this repository's Python: --run creates an isolated real daemon.
The provider wrapper only adds a durable test budget, never changes responses.
"""
from __future__ import annotations

import argparse
import asyncio
import json
import os
import socket
import sqlite3
import sys
import tempfile
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import httpx
import yaml

from aster.domain import to_dict
from aster.platform.processes import ProcessSupervisor

REPO = Path(__file__).resolve().parents[1]
REPORT = REPO / "artifacts/validation/runtime-live-20261010.json"
MAX_REQUESTS = 40  # This one authorized acceptance batch; never reset on resume.


def ledger(root):
    db = sqlite3.connect(root / "acceptance-budget.db", timeout=30)
    db.execute("CREATE TABLE IF NOT EXISTS calls(id INTEGER PRIMARY KEY, label TEXT, started REAL, elapsed REAL, status TEXT, usage TEXT, shape TEXT)")
    db.execute("CREATE TABLE IF NOT EXISTS control(key TEXT PRIMARY KEY,value TEXT)")
    db.execute("INSERT OR IGNORE INTO control VALUES('label','initial')")
    db.execute("INSERT OR IGNORE INTO control VALUES('models','0')")
    db.commit()
    return db


def label(root, value):
    with ledger(root) as db:
        db.execute("UPDATE control SET value=? WHERE key='label'", (value,))


def install_guard(root):
    from aster.providers.deepseek import DeepSeekResponsesClient
    generate = DeepSeekResponsesClient.generate
    refresh = DeepSeekResponsesClient.refresh_capabilities

    async def guarded(self, messages, tools, **kwargs):
        if self.max_output_tokens > 2048 or self.max_retries != 0:
            raise RuntimeError("Acceptance output/retry guard rejected configuration")
        shape = {"messages": len(messages), "tools": len(tools),
                 "reasoning_items": sum(item.get("type") == "reasoning" for m in messages for item in m.provider_items),
                 "image_parts": sum(p.get("type") == "input_image" for m in messages for p in m.metadata.get("content_parts", [])),
                 "summary": "Summarize the quoted archived conversation" in kwargs.get("system", "")}
        with ledger(root) as db:
            db.execute("BEGIN IMMEDIATE")
            if db.execute("SELECT count(*) FROM calls").fetchone()[0] >= MAX_REQUESTS:
                raise RuntimeError("Authorized 40 request batch is exhausted")
            current = db.execute("SELECT value FROM control WHERE key='label'").fetchone()[0]
            receipt = db.execute("INSERT INTO calls(label,started,status,shape) VALUES(?,?,'unknown',?)",
                                 (current, time.time(), json.dumps(shape))).lastrowid
        started = time.monotonic()
        try:
            turn = await generate(self, messages, tools, **kwargs)
        except BaseException as error:
            with ledger(root) as db:
                db.execute("UPDATE calls SET elapsed=?,status=? WHERE id=?", (time.monotonic()-started, type(error).__name__, receipt))
            raise
        with ledger(root) as db:
            db.execute("UPDATE calls SET elapsed=?,status='success',usage=? WHERE id=?",
                       (time.monotonic()-started, json.dumps(to_dict(turn.usage)), receipt))
        return turn

    async def metadata(self, **kwargs):
        with ledger(root) as db:
            db.execute("UPDATE control SET value=CAST(value AS INTEGER)+1 WHERE key='models'")
        return await refresh(self, **kwargs)

    DeepSeekResponsesClient.generate = guarded
    DeepSeekResponsesClient.refresh_capabilities = metadata


def prepare(root):
    from docx import Document
    workspace = root / "workspace"
    data = root / "data"
    for path in (workspace / "nested/research", workspace / "inbox", workspace / "out", data):
        path.mkdir(parents=True, exist_ok=True)
    fixture = workspace / "nested/research/settings.ini"
    if not fixture.exists():
        fixture.write_text("mode=draft\nreference=ASTER_REFERENCE_CEDAR_8421\n", encoding="utf-8")
    (workspace / "verify.py").write_text(
        "import json\nfrom pathlib import Path\np=Path('out/script-result.json')\n"
        "count=Path('out/script-count.txt')\nassert not count.exists(), 'duplicate script effect'\n"
        "assert 'mode=ready' in Path('nested/research/settings.ini').read_text()\n"
        "p.write_text(json.dumps({'sum': sum([11,17,23]),'verified':True}))\n"
        "count.write_text('1')\nprint(p.read_text())\n", encoding="utf-8")
    document = Document()
    document.add_heading("Aster source fixture", 0)
    document.add_paragraph("The measured conversion was 73.25 percent across 16 batches. Reference: ASTER_DOC_LARCH_619.")
    document.save(workspace / "source.docx")
    config = {"model": "deepseek-flash", "reasoning_effort": "low", "max_output_tokens": 2048,
              "allowed_roots": [str(workspace)], "runtime": {"max_iterations": 8, "max_model_requests": 8,
              "max_tool_calls": 20, "model_retries": 0, "model_timeout_seconds": 120,
              "deadline_seconds": 300, "reserve_tokens": 2048, "context_tokens": 60000},
              "permissions": {"auto_approve_tools": ["file_edit", "file_write", "process_exec", "browser_download"]}}
    (data / "config.yaml").write_text(yaml.safe_dump(config), encoding="utf-8")
    return data, workspace


class FixtureServer(BaseHTTPRequestHandler):
    downloads = 0

    def log_message(self, *args):
        pass

    def do_GET(self):
        if self.path == "/sample.csv":
            type(self).downloads += 1
            body = b"name,value\ncedar,47\n"
            self.send_response(200)
            self.send_header("Content-Type", "text/csv")
            self.send_header("Content-Disposition", 'attachment; filename="sample.csv"')
        else:
            body = b'<html><title>Local fixture</title><body style="background:white"><h1 style="font-size:56px">ORCHID 7318</h1><a id="download" href="/sample.csv" download>Download CSV</a></body></html>'
            self.send_response(200)
            self.send_header("Content-Type", "text/html")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


async def request(client, method, path, **kwargs):
    response = await client.request(method, path, **kwargs)
    response.raise_for_status()
    return response.json()


async def finished(client, rid):
    from aster.state.store import TERMINAL
    stop = time.monotonic() + 360
    while time.monotonic() < stop:
        run = await request(client, "GET", f"/runs/{rid}")
        if run["status"] in TERMINAL:
            return run
        await asyncio.sleep(.25)
    raise TimeoutError("Acceptance run did not terminate")


def save_report(root, cases):
    with ledger(root) as db:
        db.row_factory = sqlite3.Row
        rows = [dict(row) for row in db.execute("SELECT * FROM calls ORDER BY id")]
        metadata = int(db.execute("SELECT value FROM control WHERE key='models'").fetchone()[0])
    for row in rows:
        row["usage"] = json.loads(row["usage"]) if row["usage"] else None
        row["shape"] = json.loads(row["shape"])
    report = {"scratch": str(root), "generation_limit": MAX_REQUESTS, "generation_attempts": len(rows),
              "metadata_requests": metadata, "cases": cases, "requests": rows}
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    return report


async def acceptance(root, selected):
    data, workspace = prepare(root)
    cases = json.loads(REPORT.read_text(encoding="utf-8"))["cases"] if REPORT.exists() else []
    supervisor = ProcessSupervisor()
    cancel = asyncio.Event()
    fixture = ThreadingHTTPServer(("127.0.0.1", 0), FixtureServer)
    threading.Thread(target=fixture.serve_forever, daemon=True).start()
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    env = dict(os.environ, PYTHONPATH=str(REPO / "src"), PYTHONIOENCODING="utf-8")
    daemon = asyncio.create_task(supervisor.run([sys.executable, str(Path(__file__)), "--daemon", str(root), str(port)],
        cwd=REPO, artifact_dir=root / "daemon-logs", env=env, cancel=cancel, timeout=3600))
    try:
        async with httpx.AsyncClient(base_url=f"http://127.0.0.1:{port}", timeout=30, trust_env=False) as client:
            for _ in range(200):
                try:
                    health = await request(client, "GET", "/health")
                    if health["status"] == "ok":
                        break
                except httpx.TransportError:
                    pass
                if daemon.done():
                    raise RuntimeError("Daemon failed to start: " + (await daemon).stderr[-1000:])
                await asyncio.sleep(.1)
            else:
                raise TimeoutError("Daemon startup")
            assert (await client.get("/sessions")).status_code == 401
            client.headers["Authorization"] = "Bearer " + (data / ".api-token").read_text().strip()

            async def session():
                return (await request(client, "POST", "/sessions", json={"cwd": str(workspace)}))["id"]

            async def run_case(name, prompt, sid=None, cli=False):
                label(root, name)
                sid = sid or await session()
                started = time.monotonic()
                if cli:
                    result = await supervisor.run([sys.executable, "-m", "aster", "--data-dir", str(data), "run", prompt,
                        "--session", sid, "--json", "--no-stream", "--unattended"], cwd=REPO,
                        artifact_dir=root / "cli-logs", env=env, timeout=330)
                    value = json.loads(result.stdout)
                    rid = value["run_id"]
                else:
                    submitted = await request(client, "POST", f"/sessions/{sid}/runs", json={"text": prompt, "unattended": True})
                    rid = submitted["id"]
                run = await finished(client, rid)
                item = {"case": name, "session_id": sid, "run_id": rid, "elapsed": time.monotonic()-started,
                        "status": run["status"], "answer": run["answer"], "error": run["error"],
                        "budget": run["budget"], "usage": run["usage"], "passed": False}
                cases.append(item)
                save_report(root, cases)
                assert run["status"] == "completed", (name, run["status"], run["error"])
                return item, run

            if "files" in selected:
                item, run = await run_case("files", "Find the settings.ini file below this workspace using glob or text_search. Read it and change only mode=draft to mode=ready using file_edit and the read SHA256. Use tool_search select:file_edit to discover it. Report the reference value you read. Complete the edit exactly once.")
                assert (workspace / "nested/research/settings.ini").read_text() == "mode=ready\nreference=ASTER_REFERENCE_CEDAR_8421\n"
                assert "ASTER_REFERENCE_CEDAR_8421" in item["answer"]
                assert sum(t["tool_name"] == "file_edit" and t["state"] == "checkpointed" for t in run["tool_executions"]) == 1
                item["passed"] = True
            if "script" in selected:
                item, run = await run_case("script", f"Discover process_exec with tool_search select:process_exec. Execute exactly once the already reviewed fixture: executable={sys.executable}, argv=[\"verify.py\"], cwd={workspace}. Then use file_read on out/script-result.json to verify its sum and verified field. Do not edit the script. Report the result.", cli=True)
                assert json.loads((workspace / "out/script-result.json").read_text()) == {"sum": 51, "verified": True}
                assert (workspace / "out/script-count.txt").read_text() == "1"
                item["passed"] = True
            if "browser" in selected:
                item, run = await run_case("browser", f"Use tool_search query select:browser_open,browser_screenshot,browser_download. Call browser_open with url http://127.0.0.1:{fixture.server_port}/. Take a screenshot and read its large heading from the image. Download the CSV using browser_download selector #download exactly once. Report the visual heading and downloaded path. Do not use snapshot or HTTP text to read the heading.")
                assert "ORCHID 7318" in item["answer"]
                assert FixtureServer.downloads == 1
                files = list((data / "artifacts").rglob("*sample.csv"))
                assert len(files) == 1 and files[0].read_bytes() == b"name,value\ncedar,47\n"
                assert any(r["label"] == "browser" and r["shape"]["image_parts"] > 0 for r in save_report(root, cases)["requests"])
                item["passed"] = True
            if "document" in selected:
                item, run = await run_case("document", "Discover document_extract. Read source.docx using that tool; answer the measured conversion, batch count and reference code, citing the filename and exact paragraph number supplied by extraction. Do not infer unreported facts.")
                assert all(x in item["answer"] for x in ("73.25", "16", "ASTER_DOC_LARCH_619", "source.docx", "2"))
                item["passed"] = True
            if "history" in selected:
                item, run = await run_case("history", "Search other sessions for ASTER_REFERENCE using session_search. Report the exact reference value found and the source session ID. Do not search workspace files or invent a value.")
                source = next(c["session_id"] for c in cases if c["case"] == "files" and c["passed"])
                assert "ASTER_REFERENCE_CEDAR_8421" in item["answer"] and source in item["answer"]
                item["passed"] = True
            if "watcher" in selected:
                from docx import Document
                label(root, "watcher")
                sid = await session()
                job = await request(client, "POST", "/automations", json={"kind": "watcher", "session_id": sid,
                    "prompt": "A new document arrived. Discover document_extract and file_write. Extract inbox/incoming.docx. Write out/watcher-result.txt once with its exact reference code and paragraph number. Do not modify input files.",
                    "config": {"path": str(workspace / "inbox"), "pattern": "*.docx", "debounce_seconds": .2,
                               "stable_seconds": .2, "max_runs_per_day": 1}})
                await asyncio.sleep(2)
                document = Document()
                document.add_paragraph("New document reference is ASTER_WATCH_MAPLE_528.")
                document.save(workspace / "inbox/incoming.docx")
                started = time.monotonic()
                for _ in range(120):
                    state = await request(client, "GET", f"/sessions/{sid}")
                    if state["runs"]:
                        break
                    await asyncio.sleep(.25)
                run = await finished(client, state["runs"][0]["id"])
                await request(client, "DELETE", f"/automations/{job['id']}")
                item = {"case": "watcher", "session_id": sid, "run_id": run["id"], "elapsed": time.monotonic()-started,
                        "status": run["status"], "answer": run["answer"], "error": run["error"],
                        "budget": run["budget"], "usage": run["usage"], "passed": False}
                cases.append(item)
                assert run["status"] == "completed"
                assert "ASTER_WATCH_MAPLE_528" in (workspace / "out/watcher-result.txt").read_text()
                state = await request(client, "GET", f"/sessions/{sid}")
                assert len(state["runs"]) == 1
                item["passed"] = True
            if "compression" in selected:
                from aster.domain import Message
                from aster.engine.context import ContextManager, estimate_tokens
                from aster.state import Store
                from aster.tools.factory import ToolFactory
                from aster.tools.interaction import build_interaction_tools

                sid = next(c["session_id"] for c in cases if c["case"] == "browser" and c["passed"])
                (workspace / "padding.txt").write_text("Archive evidence row: cedar=47, larch=16.\n" * 140 + "END_MARKER=ASTER_CONTINUE_924\n", encoding="utf-8")
                padding, _ = await run_case("compression_setup", "Read padding.txt with file_read once. Report its final END_MARKER value. It is background evidence for this test; no changes are requested.", sid=sid)
                assert "ASTER_CONTINUE_924" in padding["answer"]
                padding["passed"] = True
                continuation = "Continue from prior evidence: report the visual heading previously read from the browser screenshot, and the END_MARKER from padding.txt. Do not reread files or repeat downloads."
                store = Store(data / "runtime.db")
                before = to_dict(store.history(sid))
                active = store.active_history(sid) + [Message("user", continuation)]
                overhead = store.get_checkpoint(sid)["metadata"].get("overhead_tokens", 0)
                system = store.get_run(padding["run_id"])["system"]
                factory = ToolFactory(data)
                base = {"directory_list", "glob", "text_search", "file_read", "tool_search", "use_skill", "memory_read", "session_search", "ask_user_question", "set_work_plan"}
                specs = sorted([t.spec for t in factory.tools + build_interaction_tools() if t.spec.name in base], key=lambda s: s.name)
                await factory.close()
                candidates = []
                for window in range(6144, 40000, 256):
                    manager = ContextManager(context_tokens=window, reserve_tokens=2048, retain_turns=2)
                    manager.overhead_tokens = overhead
                    try:
                        shaped = manager.compact(active, specs, system)
                    except RuntimeError:
                        continue
                    if shaped is active:
                        continue
                    older = active[:len(active)-len(shaped)+1]
                    groups, current, pending = [], [], set()
                    for message in older:
                        current.append(message)
                        pending.update(call.id for call in message.tool_calls)
                        if message.role == "tool":
                            pending.discard(message.tool_call_id)
                        if not pending:
                            groups.append(current)
                            current = []
                    if all(estimate_tokens(group) <= (window-2048)*.5 for group in groups):
                        candidates.append(window)
                assert candidates, "No safe semantic compaction window for real history"
                config = yaml.safe_load((data / "config.yaml").read_text())
                config["runtime"].update(context_tokens=max(candidates), retained_context_rounds=2)
                (data / "config.yaml").write_text(yaml.safe_dump(config), encoding="utf-8")
                item, run = await run_case("compression", continuation, sid=sid)
                assert all(value in item["answer"] for value in ("ORCHID 7318", "ASTER_CONTINUE_924"))
                events = store.events(run_id=item["run_id"], limit=10000)
                assert any(e["type"] == "context.compacted" for e in events)
                assert any(e["type"] == "context.summary.usage" for e in events)
                assert any(m.metadata.get("semantic") for m in store.active_history(sid))
                after = to_dict(store.history(sid))
                assert after[:len(before)] == before
                assert any(p.get("type") == "input_image" for m in after for p in m["metadata"].get("content_parts", []))
                assert any(p.get("type") == "reasoning" for m in after for p in m["provider_items"])
                item.update(passed=True, context_tokens=max(candidates), archive_preserved=True,
                            active_messages=len(store.active_history(sid)), raw_messages=len(after))
                store.close()
            if "restart_continuation" in selected:
                sid = next(c["session_id"] for c in cases if c["case"] == "compression" and c["passed"])
                item, run = await run_case("restart_continuation", "After this service restart, answer only the two identifiers already established in saved context: the image heading and the padding END_MARKER. Do not use tools or repeat any side effect.", sid=sid)
                assert all(value in item["answer"] for value in ("ORCHID 7318", "ASTER_CONTINUE_924"))
                assert not run["tool_executions"]
                item["passed"] = True
            report = save_report(root, cases)
            print(json.dumps({"attempts": report["generation_attempts"], "cases": [{k:c[k] for k in ("case", "status", "passed")} for c in cases]}))
    finally:
        save_report(root, cases)
        cancel.set()
        await supervisor.close()
        await daemon
        fixture.shutdown()
        fixture.server_close()


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--daemon":
        from aster.cli import main
        location = Path(sys.argv[2])
        install_guard(location)
        raise SystemExit(main(["--data-dir", str(location / "data"), "daemon", "serve", "--port", sys.argv[3]]))
    parser = argparse.ArgumentParser()
    parser.add_argument("--run", action="store_true", required=True)
    parser.add_argument("--root", type=Path)
    parser.add_argument("--cases", default="files,script,browser,document,history,watcher")
    options = parser.parse_args()
    location = options.root or Path(tempfile.mkdtemp(prefix="aster-e2e-"))
    print("Isolated acceptance workspace:", location, flush=True)
    asyncio.run(acceptance(location, set(options.cases.split(","))))
