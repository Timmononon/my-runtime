import asyncio
import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from aster.domain import ToolContext
from aster.platform.processes import ProcessSupervisor
from aster.tools.documents import DocumentTool


async def test_four_document_formats_with_source_locations(tmp_path):
    document = Document()
    document.add_heading("Heading", 1)
    document.add_paragraph("Document fact")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table value"
    document.save(tmp_path / "sample.docx")
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.active["A1"] = "Workbook fact"
    workbook.active["B2"] = "=1+1"
    workbook.save(tmp_path / "sample.xlsx")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Slide fact"
    slide.notes_slide.notes_text_frame.text = "Notes fact"
    presentation.save(tmp_path / "sample.pptx")
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 250 Td (PDF fact) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(tmp_path / "sample.pdf")
    tool = DocumentTool(ProcessSupervisor())
    context = ToolContext("s", "r", tmp_path, tmp_path / "artifacts", asyncio.Event())
    for extension, expected, key in (("docx", "Document fact", "paragraph"), ("xlsx", "Workbook fact", "cell"),
                                     ("pptx", "Slide fact", "slide"), ("pdf", "PDF fact", "page")):
        result = await tool.execute({"path": f"sample.{extension}"}, context)
        assert not result.is_error, result.content
        assert expected in result.content and key in result.content
        payload = json.loads(Path(result.artifacts[-1]["path"]).read_text(encoding="utf-8"))
        assert payload["source"].endswith(f"sample.{extension}")
        assert payload["units"]


async def test_scanned_pdf_explicit_ocr_warning(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(tmp_path / "scan.pdf")
    context = ToolContext("s", "r", tmp_path, tmp_path / "artifacts", asyncio.Event())
    result = await DocumentTool(ProcessSupervisor()).execute({"path": "scan.pdf"}, context)
    assert not result.is_error, result.content
    assert "OCR may be required" in result.content


async def test_document_budget_and_corruption_have_explicit_results(tmp_path):
    document = Document()
    for number in range(20):
        document.add_paragraph(f"Paragraph {number} with a long observation")
    document.save(tmp_path / "long.docx")
    (tmp_path / "broken.docx").write_bytes(b"not a valid Office zip file")
    supervisor = ProcessSupervisor()
    tool = DocumentTool(supervisor)
    context = ToolContext("s", "r", tmp_path, tmp_path / "artifacts", asyncio.Event())
    try:
        limited = await tool.execute({"path": "long.docx", "max_units": 2, "max_chars": 25}, context)
        assert not limited.is_error and "Extraction budget reached" in limited.content
        extracted = json.loads(Path(limited.artifacts[-1]["path"]).read_text(encoding="utf-8"))
        assert len(extracted["units"]) <= 2
        assert sum(len(item["text"]) for item in extracted["units"]) <= 25
        corrupt = await tool.execute({"path": "broken.docx"}, context)
        assert corrupt.is_error and corrupt.error_category == "document" and corrupt.outcome == "known"
    finally:
        await supervisor.close()
